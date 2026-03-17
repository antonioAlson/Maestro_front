import sys
import io

# Configurar saída UTF-8 para Windows (apenas se ainda não foi configurado)
if sys.platform == 'win32':
    try:
        if not isinstance(sys.stdout, io.TextIOWrapper) or sys.stdout.encoding != 'utf-8':
            sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
        if not isinstance(sys.stderr, io.TextIOWrapper) or sys.stderr.encoding != 'utf-8':
            sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')
    except Exception:
        pass  # Já configurado ou não é necessário

import requests
from requests.auth import HTTPBasicAuth
from dotenv import load_dotenv
import os
from docxtpl import DocxTemplate, InlineImage
from datetime import datetime
from docx.shared import Mm
import qrcode

load_dotenv()

JIRA_URL = os.getenv("JIRA_URL")
EMAIL = os.getenv("EMAIL")
API_TOKEN = os.getenv("API_TOKEN")

# Caminhos dos arquivos
TEMPLATE_PATH = os.path.join("scripts", "projetos", "default.docx")
ESPELHOS_BASE_DIR = os.path.join("src", "temp", "espelhos")
OUTPUT_DIR = os.path.join(ESPELHOS_BASE_DIR, "arquivos")
LOGS_DIR = os.path.join(ESPELHOS_BASE_DIR, "logs")
QRCODE_TEMP_DIR = os.path.join("src", "temp", "qrcodes")

# Criar diretórios de saída se não existirem
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(LOGS_DIR, exist_ok=True)
os.makedirs(QRCODE_TEMP_DIR, exist_ok=True)


def buscar_dados_card(card_id):
    """
    Busca os dados de um card específico no Jira.
    
    Args:
        card_id (str): ID do card (ex: TENSYLON-819)
    
    Returns:
        dict: Dicionário com os campos extraídos ou None se houver erro
    """
    url = f"{JIRA_URL}/rest/api/3/issue/{card_id}"
    
    headers = {
        "Accept": "application/json"
    }
    
    # Especificar apenas os campos que queremos buscar
    params = {
        "fields": "customfield_11298,customfield_11331,customfield_11071,customfield_11353,customfield_10040"
    }
    
    try:
        response = requests.get(
            url,
            headers=headers,
            params=params,
            auth=HTTPBasicAuth(EMAIL, API_TOKEN),
            timeout=30
        )
        
        if response.status_code == 200:
            data = response.json()
            fields = data.get("fields", {})
            
            # Extrair os campos desejados
            card_data = {
                "ID": card_id,
                "Veiculo - Marca/Modelo": fields.get("customfield_11298"),
                "Configurações Teto": fields.get("customfield_11331"),
                "Ano Modelo": fields.get("customfield_11071"),
                "Nº do Projeto": fields.get("customfield_11353"),
                "PEDIDO CARBON": fields.get("customfield_10040")
            }
            
            return card_data
        else:
            print(f"❌ Erro ao buscar card {card_id}: HTTP {response.status_code}")
            print(f"   Resposta: {response.text}")
            return None
            
    except Exception as e:
        print(f"❌ Erro ao buscar card {card_id}: {str(e)}")
        return None


def gerar_qrcode(card_id, card_data):
    """
    Gera um QR code com as informações do card.
    
    Args:
        card_id (str): ID do card
        card_data (dict): Dicionário com os dados do card
    
    Returns:
        str: Caminho do arquivo QR code gerado
    """
    try:
        # Montar as informações para o QR code (apenas valores, sem labels)
        qr_data = f"""{card_data.get('Veiculo - Marca/Modelo', '')}
{card_data.get('Configurações Teto', '')}
{card_data.get('Ano Modelo', '')}
{card_data.get('Nº do Projeto', '')}
{card_data.get('PEDIDO CARBON', '')}"""
        
        # Criar o QR code
        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=10,
            border=4,
        )
        qr.add_data(qr_data)
        qr.make(fit=True)
        
        # Gerar a imagem
        img = qr.make_image(fill_color="black", back_color="white")
        
        # Salvar a imagem temporariamente
        qr_filename = f"qr_{card_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.png"
        qr_path = os.path.join(QRCODE_TEMP_DIR, qr_filename)
        img.save(qr_path)
        
        return qr_path
        
    except Exception as e:
        print(f"   ⚠️  Erro ao gerar QR code: {str(e)}")
        return None


def gerar_espelho(card_id, card_data):
    """
    Gera um documento Word a partir do template default.docx.
    
    Args:
        card_id (str): ID do card
        card_data (dict): Dicionário com os dados extraídos do card
    
    Returns:
        str: Caminho do arquivo gerado ou None se houver erro
    
    Nota: O arquivo default.docx deve conter os placeholders no formato:
        {{ MODELO_VEICULO }}, {{ TIPO_TETO }}, {{ ANO_VEICULO }},
        {{ NUMERO_PROJETO }}, {{ DATA_PROJETO }}, {{ QUANTIDADE_PECAS }}, {{ NUMERO_ORDEM }}
    """
    try:
        # Verificar se o template existe
        if not os.path.exists(TEMPLATE_PATH):
            print(f"   ❌ Template não encontrado: {TEMPLATE_PATH}")
            return None
        
        # Carregar o template
        doc = DocxTemplate(TEMPLATE_PATH)
        
        # Obter data atual formatada
        data_atual = datetime.now().strftime("%d/%m/%Y")
        
        # Gerar QR code com as informações do card
        print(f"   📱 Gerando QR code...")
        qr_path = gerar_qrcode(card_id, card_data)
        
        # Preparar o contexto com os placeholders
        context = {
            'MODELO_VEICULO': card_data.get('Veiculo - Marca/Modelo', ''),
            'TIPO_TETO': card_data.get('Configurações Teto', ''),
            'ANO_VEICULO': card_data.get('Ano Modelo', ''),
            'NUMERO_PROJETO': card_data.get('Nº do Projeto', ''),
            'DATA_PROJETO': data_atual,
            'QUANTIDADE_PECAS': '',
            'NUMERO_ORDEM': card_data.get('PEDIDO CARBON', '')
        }
        
        # Adicionar QR code como imagem inline ao contexto (tamanho 35mm)
        if qr_path and os.path.exists(qr_path):
            qr_image = InlineImage(doc, qr_path, width=Mm(35))
            context['QR_CODE'] = qr_image
            print(f"   ✅ QR code adicionado ao documento (35mm)")
        else:
            context['QR_CODE'] = ''
            print(f"   ⚠️  QR code não pôde ser gerado")
        
        # Debug: mostrar os valores que serão substituídos
        print(f"   🔍 Valores para substituição:")
        for key, value in context.items():
            if key != 'QR_CODE':  # Não mostrar objeto de imagem
                print(f"      {{ {key} }} = '{value}'")
        
        # Renderizar o documento com os dados
        doc.render(context)
        
        # Gerar nome do arquivo de saída usando número do pedido, data e hora
        numero_pedido = card_data.get('PEDIDO CARBON', card_id)
        timestamp = datetime.now().strftime("%d.%m.%Y %H-%M")
        output_filename = f"{numero_pedido} {timestamp}.docx"
        output_path = os.path.join(OUTPUT_DIR, output_filename)
        
        # Salvar o documento
        doc.save(output_path)
        
        # Limpar arquivo QR code temporário
        if qr_path and os.path.exists(qr_path):
            try:
                os.remove(qr_path)
            except:
                pass  # Ignorar erros ao remover arquivo temporário
        
        return output_path
        
    except Exception as e:
        print(f"❌ Erro ao gerar espelho para {card_id}: {str(e)}")
        import traceback
        traceback.print_exc()
        return None


def processar_cards(card_ids):
    """
    Processa uma lista de IDs de cards e gera espelhos em Word.
    
    Args:
        card_ids (list): Lista de IDs dos cards
    
    Returns:
        list: Lista de dicionários com os dados de cada card e caminho do arquivo gerado
    """
    if not card_ids:
        print("⚠️  Nenhum ID de card foi informado.")
        return []
    
    # Criar arquivo de log com timestamp
    log_timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    log_file_path = os.path.join(LOGS_DIR, f"espelhos_{log_timestamp}.log")
    
    with open(log_file_path, "w", encoding="utf-8") as log_file:
        # Escrever cabeçalho do log
        log_file.write(f"{'='*70}\n")
        log_file.write(f"GERAÇÃO DE ESPELHOS - INICIADO\n")
        log_file.write(f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        log_file.write(f"{'='*70}\n\n")
        log_file.write(f"Total de cards a processar: {len(card_ids)}\n")
        log_file.write(f"IDs: {', '.join(card_ids)}\n\n")
        log_file.write(f"{'='*70}\n\n")
        
        print("\n" + "="*70)
        print("📋 CRIAR ESPELHOS - Consultando cards no Jira e gerando documentos")
        print("="*70)
        print(f"Total de cards a processar: {len(card_ids)}")
        print(f"📝 Log: {log_file_path}")
        print(f"\n⚠️  IMPORTANTE: O arquivo default.docx deve conter os placeholders:")
        print(f"   Textos: {{ MODELO_VEICULO }}, {{ TIPO_TETO }}, {{ ANO_VEICULO }},")
        print(f"           {{ NUMERO_PROJETO }}, {{ DATA_PROJETO }}, {{ QUANTIDADE_PECAS }}, {{ NUMERO_ORDEM }}")
        print(f"   QR Code: {{ QR_CODE }} (posicione onde quiser que apareça o QR code)")
        print()
        
        resultados = []
        espelhos_gerados = 0
        
        for idx, card_id in enumerate(card_ids, 1):
            log_file.write(f"[{idx}/{len(card_ids)}] Processando {card_id}...\n")
            print(f"[{idx}/{len(card_ids)}] Buscando dados de {card_id}...")
            
            card_data = buscar_dados_card(card_id)
            
            if card_data:
                print(f"   ✅ Dados extraídos com sucesso!")
                log_file.write(f"   ✅ Dados extraídos\n")
                
                # Imprimir os dados extraídos
                print(f"   📌 Veiculo - Marca/Modelo: {card_data['Veiculo - Marca/Modelo']}")
                print(f"   📌 Configurações Teto: {card_data['Configurações Teto']}")
                print(f"   📌 Ano Modelo: {card_data['Ano Modelo']}")
                print(f"   📌 Nº do Projeto: {card_data['Nº do Projeto']}")
                print(f"   📌 PEDIDO CARBON: {card_data['PEDIDO CARBON']}")
                
                log_file.write(f"   Veiculo: {card_data['Veiculo - Marca/Modelo']}\n")
                log_file.write(f"   Teto: {card_data['Configurações Teto']}\n")
                log_file.write(f"   Ano: {card_data['Ano Modelo']}\n")
                log_file.write(f"   Projeto: {card_data['Nº do Projeto']}\n")
                log_file.write(f"   PEDIDO: {card_data['PEDIDO CARBON']}\n")
                
                # Gerar o espelho em Word
                print(f"   📄 Gerando espelho...")
                output_path = gerar_espelho(card_id, card_data)
                
                if output_path:
                    print(f"   ✅ Espelho gerado: {output_path}")
                    log_file.write(f"   ✅ Espelho gerado: {os.path.basename(output_path)}\n")
                    card_data['arquivo_espelho'] = output_path
                    espelhos_gerados += 1
                else:
                    print(f"   ⚠️  Não foi possível gerar o espelho.")
                    log_file.write(f"   ⚠️  Erro ao gerar espelho\n")
                    card_data['arquivo_espelho'] = None
                
                resultados.append(card_data)
                print()
                log_file.write("\n")
            else:
                print(f"   ⚠️  Não foi possível extrair dados deste card.\n")
                log_file.write(f"   ❌ Falha ao extrair dados\n\n")
        
        # Finalizar log
        log_file.write(f"{'='*70}\n")
        log_file.write(f"✅ PROCESSAMENTO CONCLUÍDO\n")
        log_file.write(f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        log_file.write(f"{'='*70}\n")
        log_file.write(f"Cards processados: {len(resultados)}/{len(card_ids)}\n")
        log_file.write(f"Espelhos gerados: {espelhos_gerados}\n")
        log_file.write(f"Pasta de saída: {OUTPUT_DIR}\n")
        
        print("="*70)
        print(f"✅ Processamento concluído!")
        print(f"   📊 {len(resultados)}/{len(card_ids)} cards processados com sucesso")
        print(f"   📄 {espelhos_gerados} espelhos gerados")
        print(f"   📁 Arquivos salvos em: {OUTPUT_DIR}")
        print(f"   📝 Log salvo em: {log_file_path}")
        print("="*70 + "\n")
    
    return resultados
